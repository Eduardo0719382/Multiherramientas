"""Pestaña "Convertidor universal": video, audio, imagen, PDF, Word, Excel/CSV, PowerPoint, texto.

Motores:
  - ffmpeg (en PATH)            -> video y audio
  - Pillow / img2pdf / PyMuPDF  -> imágenes y PDF
  - pdf2docx                    -> PDF a Word
  - python-docx / pandas        -> Word a texto, Excel/CSV
  - Microsoft Office (COM)      -> Word/Excel/PowerPoint a PDF (si está instalado)
"""
import os
import re
import shutil
import subprocess
import tempfile
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

VIDEO_IN = {".mp4", ".mkv", ".avi", ".mov", ".webm", ".flv", ".wmv", ".m4v", ".mpg", ".mpeg", ".3gp", ".ts", ".gif"}
AUDIO_IN = {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac", ".opus", ".wma", ".aiff", ".amr"}
IMAGE_IN = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".ico", ".heic", ".avif"}
PDF_IN = {".pdf"}
WORD_IN = {".docx", ".doc", ".rtf", ".odt"}
SHEET_IN = {".xlsx", ".xls", ".csv", ".ods"}
SLIDE_IN = {".pptx", ".ppt", ".odp"}
TEXT_IN = {".txt", ".md", ".html", ".htm"}

# categoría -> formatos de salida
OUTPUTS = {
    "Video": ["MP4", "MKV", "AVI", "MOV", "WEBM", "GIF",
              "MP3 (solo audio)", "WAV (solo audio)", "M4A (solo audio)", "FLAC (solo audio)",
              "PNG (secuencia de fotogramas)"],
    "Audio": ["MP3", "WAV", "FLAC", "OGG", "M4A", "AAC", "OPUS", "WMA"],
    "Imagen": ["PNG", "JPG", "WEBP", "BMP", "TIFF", "GIF", "ICO", "PDF"],
    "PDF": ["DOCX (Word)", "TXT", "PNG (una por página)", "JPG (una por página)"],
    "Word": ["PDF", "TXT", "DOCX"],
    "Hoja de cálculo": ["XLSX", "CSV", "PDF", "HTML"],
    "Presentación": ["PDF", "PPTX"],
    "Texto": ["PDF", "TXT", "DOCX", "HTML"],
}


def categoria(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".gif":
        return "Imagen"
    for nombre, exts in (("Video", VIDEO_IN), ("Audio", AUDIO_IN), ("Imagen", IMAGE_IN), ("PDF", PDF_IN),
                         ("Word", WORD_IN), ("Hoja de cálculo", SHEET_IN), ("Presentación", SLIDE_IN),
                         ("Texto", TEXT_IN)):
        if ext in exts:
            return nombre
    return None


ALL_EXTS = sorted(VIDEO_IN | AUDIO_IN | IMAGE_IN | PDF_IN | WORD_IN | SHEET_IN | SLIDE_IN | TEXT_IN)

AUDIO_CODECS = {
    "mp3": ["-c:a", "libmp3lame", "-q:a", "2"],
    "wav": ["-c:a", "pcm_s16le"],
    "flac": ["-c:a", "flac"],
    "ogg": ["-c:a", "libvorbis", "-q:a", "5"],
    "m4a": ["-c:a", "aac", "-b:a", "192k"],
    "aac": ["-c:a", "aac", "-b:a", "192k"],
    "opus": ["-c:a", "libopus", "-b:a", "128k"],
    "wma": ["-c:a", "wmav2", "-b:a", "192k"],
}

VIDEO_CODECS = {
    "mp4": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", "20", "-c:a", "aac", "-movflags", "+faststart"],
    "mkv": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", "20", "-c:a", "aac"],
    "avi": ["-c:v", "mpeg4", "-q:v", "3", "-c:a", "libmp3lame"],
    "mov": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", "20", "-c:a", "aac"],
    "webm": ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-c:a", "libopus"],
}


def colores_tema():
    """(fondo, texto) del tema sv_ttk activo; ttk.Style().lookup no los devuelve con ese tema."""
    try:
        import sv_ttk
        if sv_ttk.get_theme() == "dark":
            return "#1c1c1c", "#fafafa"
    except Exception:
        pass
    return "#fafafa", "#1c1c1c"


class ConvertError(Exception):
    pass


def _ffmpeg(args, cancel_check=None):
    if not shutil.which("ffmpeg"):
        raise ConvertError("No se encontró ffmpeg en el PATH (necesario para video/audio).")
    r = subprocess.run(["ffmpeg", "-y", "-loglevel", "error"] + args, capture_output=True,
                       creationflags=NO_WINDOW)
    if r.returncode != 0:
        msg = r.stderr.decode("utf-8", "replace").strip()
        if "does not contain any stream" in msg:
            raise ConvertError("el archivo no contiene audio.")
        raise ConvertError("ffmpeg: " + msg[-400:])


def _office_a_pdf(src, dst):
    """Word / Excel / PowerPoint -> PDF mediante COM (requiere Office instalado)."""
    ext = os.path.splitext(src)[1].lower()
    src, dst = os.path.abspath(src), os.path.abspath(dst)
    if ext in WORD_IN:
        script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject Word.Application
$app.Visible = $false; $app.DisplayAlerts = 0
try {{ $d = $app.Documents.Open('{src}', $false, $true); $d.ExportAsFixedFormat('{dst}', 17); $d.Close($false) }}
finally {{ $app.Quit() }}
"""
    elif ext in SHEET_IN:
        script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject Excel.Application
$app.Visible = $false; $app.DisplayAlerts = $false
try {{ $w = $app.Workbooks.Open('{src}', 0, $true); $w.ExportAsFixedFormat(0, '{dst}'); $w.Close($false) }}
finally {{ $app.Quit() }}
"""
    elif ext in SLIDE_IN:
        script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject PowerPoint.Application
try {{ $p = $app.Presentations.Open('{src}', $true, $false, $false); $p.SaveAs('{dst}', 32); $p.Close() }}
finally {{ $app.Quit() }}
"""
    else:
        raise ConvertError("Formato no soportado para PDF vía Office.")
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       capture_output=True, creationflags=NO_WINDOW)
    if r.returncode != 0 or not os.path.exists(dst):
        raise ConvertError("Microsoft Office no pudo convertir el archivo (¿está instalado?): "
                           + r.stderr.decode("utf-8", "replace").strip()[-300:])


def _office_save_as(src, dst, kind):
    """Convierte con Word/PowerPoint a otro formato (docx desde doc/rtf/odt, pptx desde ppt/odp)."""
    src, dst = os.path.abspath(src), os.path.abspath(dst)
    if kind == "word":
        script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject Word.Application
$app.Visible = $false; $app.DisplayAlerts = 0
try {{ $d = $app.Documents.Open('{src}', $false, $true); $d.SaveAs2('{dst}', 16); $d.Close($false) }}
finally {{ $app.Quit() }}
"""
    else:
        script = f"""
$ErrorActionPreference='Stop'
$app = New-Object -ComObject PowerPoint.Application
try {{ $p = $app.Presentations.Open('{src}', $true, $false, $false); $p.SaveAs('{dst}', 24); $p.Close() }}
finally {{ $app.Quit() }}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       capture_output=True, creationflags=NO_WINDOW)
    if r.returncode != 0 or not os.path.exists(dst):
        raise ConvertError("Microsoft Office no pudo convertir el archivo: "
                           + r.stderr.decode("utf-8", "replace").strip()[-300:])


def _texto_a_pdf(texto, dst):
    import fitz
    ancho, alto, margen, tam = 595, 842, 56, 11
    max_w = ancho - 2 * margen
    por_pagina = int((alto - 2 * margen) // (tam * 1.35))
    lineas = []
    for parrafo in texto.replace("	", "    ").splitlines() or [""]:
        actual = ""
        for palabra in parrafo.split(" "):
            prueba = palabra if not actual else actual + " " + palabra
            if fitz.get_text_length(prueba, fontname="helv", fontsize=tam) <= max_w:
                actual = prueba
            else:
                if actual:
                    lineas.append(actual)
                actual = palabra
        lineas.append(actual)
    doc = fitz.open()
    for i in range(0, max(len(lineas), 1), por_pagina):
        page = doc.new_page(width=ancho, height=alto)
        y = margen + tam
        for ln in lineas[i:i + por_pagina]:
            page.insert_text((margen, y), ln, fontsize=tam, fontname="helv")
            y += tam * 1.35
    doc.save(dst)
    doc.close()


def _leer_texto(src):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(src, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return ""


def _word_a_texto(src):
    ext = os.path.splitext(src)[1].lower()
    if ext == ".docx":
        import docx
        d = docx.Document(src)
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            for fila in t.rows:
                partes.append("\t".join(c.text for c in fila.cells))
        return "\n".join(partes)
    tmp = os.path.join(tempfile.mkdtemp(), "t.docx")
    _office_save_as(src, tmp, "word")
    return _word_a_texto(tmp)


def convertir(src, fmt, carpeta, cancel_check=lambda: False):
    """Convierte `src` al formato `fmt`. Devuelve la lista de rutas creadas."""
    cat = categoria(src)
    base = os.path.splitext(os.path.basename(src))[0]
    ext_in = os.path.splitext(src)[1].lower()
    f = fmt.split(" ")[0].lower()  # "MP3 (solo audio)" -> "mp3"

    def salida(ext):
        out = os.path.join(carpeta, f"{base}.{ext}")
        if os.path.abspath(out) == os.path.abspath(src):
            out = os.path.join(carpeta, f"{base}_convertido.{ext}")
        return out

    # ----- VIDEO -----
    if cat == "Video":
        if "solo audio" in fmt:
            out = salida(f)
            _ffmpeg(["-i", src, "-vn"] + AUDIO_CODECS[f] + [out])
            return [out]
        if f == "gif":
            out = salida("gif")
            _ffmpeg(["-i", src, "-vf", "fps=12,scale=640:-1:flags=lanczos", "-loop", "0", out])
            return [out]
        if "fotogramas" in fmt:
            d = os.path.join(carpeta, f"{base}_fotogramas")
            os.makedirs(d, exist_ok=True)
            _ffmpeg(["-i", src, "-vf", "fps=1", os.path.join(d, "frame_%05d.png")])
            return [d]
        out = salida(f)
        _ffmpeg(["-i", src] + VIDEO_CODECS[f] + [out])
        return [out]

    # ----- AUDIO -----
    if cat == "Audio":
        out = salida(f)
        _ffmpeg(["-i", src, "-vn"] + AUDIO_CODECS[f] + [out])
        return [out]

    # ----- IMAGEN -----
    if cat == "Imagen":
        if f == "pdf":
            from PIL import Image
            out = salida("pdf")
            with Image.open(src) as im:
                im.convert("RGB").save(out, "PDF", resolution=150.0)
            return [out]
        if ext_in in {".heic", ".avif"}:
            # Pillow estándar no lo lee: ffmpeg como alternativa
            out = salida({"jpg": "jpg"}.get(f, f))
            _ffmpeg(["-i", src, out])
            return [out]
        from PIL import Image
        out = salida("jpg" if f == "jpg" else f)
        with Image.open(src) as im:
            if f == "ico":
                im.convert("RGBA").save(out, sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
            elif f in {"jpg", "bmp"}:
                if im.mode in ("RGBA", "LA", "P"):
                    im = im.convert("RGBA")
                    fondo = Image.new("RGB", im.size, (255, 255, 255))
                    fondo.paste(im, mask=im.split()[-1])
                    im = fondo
                else:
                    im = im.convert("RGB")
                im.save(out, "JPEG" if f == "jpg" else "BMP", **({"quality": 92} if f == "jpg" else {}))
            elif f == "gif":
                im.save(out, "GIF")
            elif f == "tiff":
                im.save(out, "TIFF")
            else:
                im.save(out, f.upper())
        return [out]

    # ----- PDF -----
    if cat == "PDF":
        if f == "docx":
            from pdf2docx import Converter
            out = salida("docx")
            cv = Converter(src)
            try:
                cv.convert(out)
            finally:
                cv.close()
            return [out]
        import fitz
        doc = fitz.open(src)
        try:
            if f == "txt":
                out = salida("txt")
                with open(out, "w", encoding="utf-8") as fh:
                    for i, p in enumerate(doc, 1):
                        fh.write(p.get_text())
                        fh.write(f"\n\n--- Página {i} ---\n\n")
                return [out]
            if f in {"png", "jpg"}:
                d = os.path.join(carpeta, f"{base}_paginas")
                os.makedirs(d, exist_ok=True)
                creados = []
                for i, p in enumerate(doc, 1):
                    if cancel_check():
                        break
                    o = os.path.join(d, f"pagina_{i:03d}.{f}")
                    p.get_pixmap(dpi=150).save(o)
                    creados.append(o)
                return [d]
        finally:
            doc.close()

    # ----- WORD -----
    if cat == "Word":
        if f == "pdf":
            out = salida("pdf")
            _office_a_pdf(src, out)
            return [out]
        if f == "txt":
            out = salida("txt")
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(_word_a_texto(src))
            return [out]
        if f == "docx":
            out = salida("docx")
            _office_save_as(src, out, "word")
            return [out]

    # ----- HOJAS DE CÁLCULO -----
    if cat == "Hoja de cálculo":
        import pandas as pd
        if f == "pdf":
            out = salida("pdf")
            _office_a_pdf(src, out)
            return [out]
        df = pd.read_csv(src, sep=None, engine="python", encoding_errors="replace") if ext_in == ".csv" \
            else pd.read_excel(src, sheet_name=None)
        if f == "csv":
            if isinstance(df, dict):
                creados = []
                for nombre, hoja in df.items():
                    o = os.path.join(carpeta, f"{base}_{re.sub(r'[^\w-]+', '_', str(nombre))}.csv") \
                        if len(df) > 1 else salida("csv")
                    hoja.to_csv(o, index=False, encoding="utf-8-sig")
                    creados.append(o)
                return creados
            out = salida("csv")
            df.to_csv(out, index=False, encoding="utf-8-sig")
            return [out]
        if f == "xlsx":
            out = salida("xlsx")
            with pd.ExcelWriter(out) as w:
                if isinstance(df, dict):
                    for nombre, hoja in df.items():
                        hoja.to_excel(w, sheet_name=str(nombre)[:31], index=False)
                else:
                    df.to_excel(w, index=False)
            return [out]
        if f == "html":
            out = salida("html")
            partes = df.values() if isinstance(df, dict) else [df]
            with open(out, "w", encoding="utf-8") as fh:
                fh.write("<meta charset='utf-8'>" + "<hr>".join(h.to_html(index=False) for h in partes))
            return [out]

    # ----- PRESENTACIÓN -----
    if cat == "Presentación":
        if f == "pdf":
            out = salida("pdf")
            _office_a_pdf(src, out)
            return [out]
        if f == "pptx":
            out = salida("pptx")
            _office_save_as(src, out, "ppt")
            return [out]

    # ----- TEXTO -----
    if cat == "Texto":
        texto = _leer_texto(src)
        if ext_in in {".html", ".htm"}:
            texto = re.sub(r"<(script|style).*?</\1>", "", texto, flags=re.S | re.I)
            texto = re.sub(r"<[^>]+>", "", texto)
        if f == "pdf":
            out = salida("pdf")
            _texto_a_pdf(texto, out)
            return [out]
        if f == "txt":
            out = salida("txt")
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(texto)
            return [out]
        if f == "docx":
            import docx
            out = salida("docx")
            d = docx.Document()
            for linea in texto.splitlines():
                d.add_paragraph(linea)
            d.save(out)
            return [out]
        if f == "html":
            import html as _h
            out = salida("html")
            with open(out, "w", encoding="utf-8") as fh:
                fh.write("<meta charset='utf-8'><pre style='font-family:sans-serif;white-space:pre-wrap'>"
                         + _h.escape(texto) + "</pre>")
            return [out]

    raise ConvertError(f"Conversión no soportada: {cat} → {fmt}")


def unir_imagenes_pdf(archivos, destino):
    import img2pdf
    from PIL import Image
    paginas = []
    tmp = tempfile.mkdtemp()
    for i, a in enumerate(archivos):
        with Image.open(a) as im:
            p = os.path.join(tmp, f"{i}.jpg")
            im.convert("RGB").save(p, "JPEG", quality=95)
            paginas.append(p)
    with open(destino, "wb") as fh:
        fh.write(img2pdf.convert(paginas))
    shutil.rmtree(tmp, ignore_errors=True)


class UniversalConverter(ttk.Frame):
    def __init__(self, master):
        super().__init__(master)
        self.archivos = []
        self.cancel = False
        self.trabajando = False
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self, padding=10)
        top.pack(fill=tk.BOTH, expand=True)

        ttk.Label(top, text="Convertidor universal: video, audio, imagen, PDF, Word, Excel/CSV, PowerPoint y texto",
                  font=("Segoe UI", 10, "bold")).pack(anchor=tk.W)

        fila = ttk.Frame(top)
        fila.pack(fill=tk.BOTH, expand=True, pady=6)

        lista_box = ttk.Frame(fila)
        lista_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        fondo, texto = colores_tema()
        self.lista = tk.Listbox(lista_box, selectmode=tk.EXTENDED, activestyle="none", bg=fondo, fg=texto,
                                relief=tk.FLAT, highlightthickness=0)
        sb = ttk.Scrollbar(lista_box, command=self.lista.yview)
        self.lista.config(yscrollcommand=sb.set)
        self.lista.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb.pack(side=tk.LEFT, fill=tk.Y)

        botones = ttk.Frame(fila, padding=(8, 0))
        botones.pack(side=tk.LEFT, fill=tk.Y)
        ttk.Button(botones, text="➕ Agregar archivos", command=self.agregar).pack(fill=tk.X, pady=2)
        ttk.Button(botones, text="➖ Quitar", command=self.quitar).pack(fill=tk.X, pady=2)
        ttk.Button(botones, text="🗑 Vaciar", command=self.vaciar).pack(fill=tk.X, pady=2)

        self.lbl_cat = ttk.Label(top, text="Arrastra archivos aquí o usa «Agregar archivos»", font=("Segoe UI", 9, "italic"))
        self.lbl_cat.pack(anchor=tk.W)

        opts = ttk.Frame(top)
        opts.pack(fill=tk.X, pady=6)
        ttk.Label(opts, text="Convertir a:").grid(row=0, column=0, sticky=tk.W)
        self.combo_fmt = ttk.Combobox(opts, state="readonly", width=34)
        self.combo_fmt.grid(row=0, column=1, sticky=tk.W, padx=6)

        ttk.Label(opts, text="Carpeta de salida:").grid(row=1, column=0, sticky=tk.W, pady=4)
        self.var_salida = tk.StringVar()
        ttk.Entry(opts, textvariable=self.var_salida, width=60).grid(row=1, column=1, sticky=tk.EW, padx=6)
        ttk.Button(opts, text="Elegir…", command=self.elegir_salida).grid(row=1, column=2)
        self.var_junto = tk.BooleanVar(value=False)
        self.chk_junto = ttk.Checkbutton(opts, text="Unir todas las imágenes en un solo PDF", variable=self.var_junto)
        self.chk_junto.grid(row=2, column=1, sticky=tk.W, padx=6)
        opts.columnconfigure(1, weight=1)
        self.combo_fmt.bind("<<ComboboxSelected>>", lambda e: self._actualizar_opciones())

        acciones = ttk.Frame(top)
        acciones.pack(fill=tk.X)
        self.btn_ir = ttk.Button(acciones, text="🔄 Convertir", style="Accent.TButton", command=self.iniciar)
        self.btn_ir.pack(side=tk.LEFT)
        self.btn_cancel = ttk.Button(acciones, text="Cancelar", command=self.cancelar, state=tk.DISABLED)
        self.btn_cancel.pack(side=tk.LEFT, padx=6)
        self.progress = ttk.Progressbar(acciones, mode="determinate")
        self.progress.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=6)

        self.log = tk.Text(top, height=8, state=tk.DISABLED, wrap=tk.WORD, bg=fondo, fg=texto,
                           relief=tk.FLAT, highlightthickness=0)
        self.log.pack(fill=tk.X, pady=(6, 0))
        self._actualizar_opciones()

    # ---------- lista ----------
    def agregar(self):
        tipos = [("Todos los soportados", " ".join("*" + e for e in ALL_EXTS)), ("Todos", "*.*")]
        self.recibir_archivos(filedialog.askopenfilenames(filetypes=tipos))

    def recibir_archivos(self, nuevos):
        for p in nuevos:
            if p not in self.archivos:
                if categoria(p) is None:
                    self._log(f"Omitido (formato no soportado): {os.path.basename(p)}")
                    continue
                self.archivos.append(p)
        self._refrescar()

    def quitar(self):
        for i in reversed(self.lista.curselection()):
            del self.archivos[i]
        self._refrescar()

    def vaciar(self):
        self.archivos.clear()
        self._refrescar()

    def _refrescar(self):
        self.lista.delete(0, tk.END)
        for p in self.archivos:
            self.lista.insert(tk.END, f"[{categoria(p)}]  {p}")
        cats = {categoria(p) for p in self.archivos}
        if not cats:
            self.lbl_cat.config(text="Arrastra archivos aquí o usa «Agregar archivos»")
            self.combo_fmt.config(values=[])
            self.combo_fmt.set("")
        elif len(cats) > 1:
            self.lbl_cat.config(text="Hay tipos mezclados (" + ", ".join(sorted(cats)) +
                                "). Convierte un tipo a la vez.")
            self.combo_fmt.config(values=[])
            self.combo_fmt.set("")
        else:
            c = cats.pop()
            self.lbl_cat.config(text=f"Tipo detectado: {c}")
            self.combo_fmt.config(values=OUTPUTS[c])
            self.combo_fmt.set(OUTPUTS[c][0])
            if not self.var_salida.get():
                self.var_salida.set(os.path.dirname(self.archivos[0]))
        self._actualizar_opciones()

    def _actualizar_opciones(self):
        cats = {categoria(p) for p in self.archivos}
        mostrar = cats == {"Imagen"} and self.combo_fmt.get() == "PDF" and len(self.archivos) > 1
        self.chk_junto.state(["!disabled"] if mostrar else ["disabled"])
        if not mostrar:
            self.var_junto.set(False)

    def elegir_salida(self):
        d = filedialog.askdirectory()
        if d:
            self.var_salida.set(d)

    # ---------- ejecución ----------
    def _log(self, txt):
        self.log.config(state=tk.NORMAL)
        self.log.insert(tk.END, txt + "\n")
        self.log.see(tk.END)
        self.log.config(state=tk.DISABLED)

    def _ui(self, fn):
        self.after(0, fn)

    def cancelar(self):
        self.cancel = True

    def iniciar(self):
        if self.trabajando:
            return
        if not self.archivos or not self.combo_fmt.get():
            messagebox.showwarning("Convertidor", "Agrega archivos del mismo tipo y elige un formato de salida.")
            return
        carpeta = self.var_salida.get().strip()
        if not carpeta:
            messagebox.showwarning("Convertidor", "Elige la carpeta de salida.")
            return
        os.makedirs(carpeta, exist_ok=True)
        self.cancel = False
        self.trabajando = True
        self.btn_ir.config(state=tk.DISABLED)
        self.btn_cancel.config(state=tk.NORMAL)
        self.progress.config(value=0)
        threading.Thread(target=self._trabajo, args=(list(self.archivos), self.combo_fmt.get(), carpeta,
                                                     self.var_junto.get()), daemon=True).start()

    def _trabajo(self, archivos, fmt, carpeta, juntar):
        ok, fallos = 0, 0
        try:
            if juntar:
                destino = os.path.join(carpeta, "imagenes_unidas.pdf")
                self._ui(lambda: self._log("Uniendo imágenes en un solo PDF..."))
                try:
                    unir_imagenes_pdf(archivos, destino)
                    ok = 1
                    self._ui(lambda: self._log(f"✔ {destino}"))
                except Exception as e:
                    fallos = 1
                    self._ui(lambda e=e: self._log(f"✘ Error: {e}"))
            else:
                for i, src in enumerate(archivos, 1):
                    if self.cancel:
                        self._ui(lambda: self._log("Cancelado."))
                        break
                    nombre = os.path.basename(src)
                    self._ui(lambda n=nombre, i=i: self._log(f"({i}/{len(archivos)}) Convirtiendo {n}..."))
                    try:
                        salidas = convertir(src, fmt, carpeta, lambda: self.cancel)
                        ok += 1
                        self._ui(lambda s=salidas: self._log("   ✔ " + ", ".join(s)))
                    except Exception as e:
                        fallos += 1
                        self._ui(lambda e=e, n=nombre: self._log(f"   ✘ {n}: {e}"))
                    self._ui(lambda v=i / len(archivos) * 100: self.progress.config(value=v))
        finally:
            def fin():
                self.trabajando = False
                self.btn_ir.config(state=tk.NORMAL)
                self.btn_cancel.config(state=tk.DISABLED)
                self.progress.config(value=100 if not fallos else self.progress["value"])
                self._log(f"Terminado: {ok} correcto(s), {fallos} con error.")
                if ok and not fallos:
                    messagebox.showinfo("Convertidor", f"Conversión completada ({ok}).\n{carpeta}")
            self._ui(fin)
